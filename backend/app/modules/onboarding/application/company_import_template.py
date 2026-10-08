"""The Excel template for bulk company import.

Three sheets:

* **Instructions** — how to fill the file in, and one line per column: whether a
  value is needed, what to enter, an example and the allowed values.
* **Companies** — the header row the importer expects (``TEMPLATE_COLUMNS``),
  dropdowns for country and source, and identifier columns formatted as text so
  Excel keeps leading zeros and long numbers as typed.
* **Lists** (hidden) — the values the dropdowns offer.

Everything a rule depends on comes from the importer itself — its columns, its
sources and their labels, its row limit — so the template cannot drift from what
the import accepts. ``COLUMN_GUIDE`` must describe every template column; a test
holds it to that.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.modules.onboarding.application.company_import_service import (
    COMPANIES_SHEET,
    IMPORT_SOURCE_LABELS,
    MAX_ROWS,
    TEMPLATE_COLUMNS,
)
from app.modules.onboarding.domain.company_identity import REGISTRATION_NUMBER_MAX
from app.modules.onboarding.domain.countries import COUNTRY_NAMES

TEMPLATE_XLSX_NAME = "company-import-template.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_LISTS_SHEET = "Lists"
#: Columns whose values are identifiers: kept as text, never as numbers.
_TEXT_COLUMNS = ("pan", "gstins", "iec", "cin", "registration_number")


@dataclass(frozen=True)
class ColumnGuide:
    column: str
    needed: str
    what: str
    example: str
    allowed: str = ""


COLUMN_GUIDE: tuple[ColumnGuide, ...] = (
    ColumnGuide(
        "name", "Yes",
        "The company's registered name, up to 255 characters.",
        "Shree Textiles Private Limited",
    ),
    ColumnGuide(
        "country", "Yes",
        "Pick the country from the list. You can also type its two-letter ISO code.",
        "India (or IN)",
        "Any country in the list",
    ),
    ColumnGuide(
        "pan", "Recommended for Indian companies",
        "10 characters: 5 letters, 4 digits, 1 letter. A row whose PAN is already in "
        "the CRM is matched to that company, not added again.",
        "ABCDE1234F",
    ),
    ColumnGuide(
        "gstins", "No",
        "One or more 15-character GSTINs, separated by a semicolon (;). Each must "
        "contain the company's PAN (characters 3 to 12).",
        "27ABCDE1234F1Z5;29ABCDE1234F1Z3",
    ),
    ColumnGuide(
        "iec", "No",
        "Importer-Exporter Code: 10 letters or digits.",
        "0512345678",
    ),
    ColumnGuide(
        "cin", "No",
        "Corporate Identity Number: 21 characters — L or U, 5 digits, 2-letter state, "
        "4-digit year, 3 letters, 6 digits.",
        "U17110MH2010PTC123456",
    ),
    ColumnGuide(
        "source", "No (blank means Manual entry)",
        "How the relationship started. Pick from the list.",
        "Referral",
        ", ".join(IMPORT_SOURCE_LABELS.values()),
    ),
    ColumnGuide(
        "industry", "No",
        "The company's industry, in your own words, up to 255 characters.",
        "Textiles",
    ),
    ColumnGuide(
        "registration_number", "Yes for a company outside India that has no PAN",
        "The number the company's own registrar issued, written as the registrar "
        f"writes it, up to {REGISTRATION_NUMBER_MAX} characters.",
        "HRB 123456",
    ),
)

_HOW_TO = (
    "Fill in the Companies sheet, one company per row, and upload the file on the "
    "Import companies page.",
    "Keep the header row as it is: do not rename, add or remove columns. A column "
    "you do not need can stay empty.",
    f"Up to {MAX_ROWS:,} companies per file. Blank rows are skipped.",
    "Every row is checked by the same rules as adding a company by hand. New companies "
    "start as leads.",
    "A row that looks like an existing company but has no PAN to confirm it is not "
    "added: it is listed for a person to decide.",
    "After the import, every row is reported: created, matched, possible duplicate or "
    "failed, with the reason.",
)

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
_WRAP = Alignment(wrap_text=True, vertical="top")


def build_template_xlsx() -> bytes:
    """The template workbook, as the bytes of an ``.xlsx`` file."""
    workbook = Workbook()
    _instructions(workbook.active)
    companies = workbook.create_sheet(COMPANIES_SHEET)
    lists = workbook.create_sheet(_LISTS_SHEET)
    country_range, source_range = _lists(lists)
    _companies(companies, country_range, source_range)
    lists.sheet_state = "hidden"
    workbook.active = 1  # open on the Companies sheet
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _instructions(sheet) -> None:
    sheet.title = "Instructions"
    sheet["A1"] = "Import companies — how to fill in this file"
    sheet["A1"].font = _TITLE
    row = 3
    for line in _HOW_TO:
        sheet.cell(row=row, column=1, value=f"• {line}")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
    row += 1
    for index, title in enumerate(
        ("Column", "Needed?", "What to enter", "Example", "Allowed values"), start=1
    ):
        cell = sheet.cell(row=row, column=index, value=title)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    for guide in COLUMN_GUIDE:
        row += 1
        for index, value in enumerate(
            (guide.column, guide.needed, guide.what, guide.example, guide.allowed), start=1
        ):
            sheet.cell(row=row, column=index, value=value).alignment = _WRAP
    for index, width in enumerate((22, 26, 70, 34, 40), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _lists(sheet) -> tuple[str, str]:
    """Fill the hidden sheet the dropdowns read; return their two ranges."""
    names = sorted(COUNTRY_NAMES.values(), key=str.casefold)
    for row, name in enumerate(names, start=1):
        sheet.cell(row=row, column=1, value=name)
    labels = list(IMPORT_SOURCE_LABELS.values())
    for row, label in enumerate(labels, start=1):
        sheet.cell(row=row, column=2, value=label)
    return (
        f"{_LISTS_SHEET}!$A$1:$A${len(names)}",
        f"{_LISTS_SHEET}!$B$1:$B${len(labels)}",
    )


def _companies(sheet, country_range: str, source_range: str) -> None:
    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[get_column_letter(index)].width = max(16, len(column) + 4)
    sheet.freeze_panes = "A2"
    last_row = MAX_ROWS + 1

    def column_range(column: str) -> str:
        letter = get_column_letter(TEMPLATE_COLUMNS.index(column) + 1)
        return f"{letter}2:{letter}{last_row}"

    # A warning rather than a stop: typing an ISO code instead of picking a name
    # is allowed, and the importer reads either.
    country = DataValidation(
        type="list", formula1=f"={country_range}", allow_blank=True,
        showErrorMessage=True, errorStyle="warning",
        errorTitle="Country", error="Pick a country from the list, or type its two-letter code.",
    )
    country.add(column_range("country"))
    source = DataValidation(
        type="list", formula1=f"={source_range}", allow_blank=True,
        showErrorMessage=True, errorStyle="stop",
        errorTitle="Source", error="Pick a source from the list.",
    )
    source.add(column_range("source"))
    sheet.add_data_validation(country)
    sheet.add_data_validation(source)

    for column in _TEXT_COLUMNS:
        index = TEMPLATE_COLUMNS.index(column) + 1
        for row in range(2, last_row + 1):
            sheet.cell(row=row, column=index).number_format = "@"


__all__ = [
    "COLUMN_GUIDE",
    "TEMPLATE_XLSX_NAME",
    "XLSX_MEDIA_TYPE",
    "build_template_xlsx",
]
