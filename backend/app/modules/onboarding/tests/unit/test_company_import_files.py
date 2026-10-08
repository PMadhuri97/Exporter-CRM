"""Reading import files (CSV and Excel), the preview, and the Excel template.

No database: these are the steps before any row is judged. Row judging and saving
are covered in ``tests/integration/test_company_import.py``.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.onboarding.application.company_import_service import (
    COMPANIES_SHEET,
    IMPORT_SOURCE_LABELS,
    TEMPLATE_COLUMNS,
    _source_for,
    preview_file,
    read_file,
)
from app.modules.onboarding.application.company_import_template import (
    COLUMN_GUIDE,
    build_template_xlsx,
)
from app.modules.onboarding.domain.countries import COUNTRY_NAMES, country_code_for
from app.modules.onboarding.domain.entities.exporter_enums import ExporterSource
from app.shared.exceptions import ValidationError

HEADER = ",".join(TEMPLATE_COLUMNS)


def _xlsx(rows: list[list], *, sheet: str = COMPANIES_SHEET, extra_sheet: str | None = None) -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    if extra_sheet:
        workbook.create_sheet(extra_sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── CSV ───────────────────────────────────────────────────────────────────────


def test_a_csv_keeps_quoted_commas_and_empty_cells_in_their_columns():
    content = (HEADER + '\n"Acme, Ltd",IN,,27ABCDE1234F1Z5,,,,Textiles,\n').encode()
    parsed = read_file("companies.csv", content)
    [(line, raw)] = parsed.rows
    assert line == 2
    assert raw["name"] == "Acme, Ltd"
    assert raw["pan"] == ""
    assert raw["gstins"] == "27ABCDE1234F1Z5"
    assert raw["industry"] == "Textiles"


def test_a_csv_with_a_byte_order_mark_reads_its_first_column():
    content = ("﻿" + HEADER + "\nAcme,IN,,,,,,,\n").encode("utf-8")
    assert read_file("companies.csv", content).header[0] == "name"


# ── Excel ─────────────────────────────────────────────────────────────────────


def test_an_excel_file_is_read_like_a_csv_with_sheet_row_numbers():
    content = _xlsx([
        list(TEMPLATE_COLUMNS),
        ["Acme", "India", None, None, 512345678, None, "Referral", "Textiles", None],
        [None] * len(TEMPLATE_COLUMNS),  # a blank row is skipped
        ["Beta", "IN", "ABCDE1234F"],
    ])
    parsed = read_file("companies.xlsx", content)
    assert parsed.header == list(TEMPLATE_COLUMNS)
    assert [line for line, _ in parsed.rows] == [2, 4]
    first = parsed.rows[0][1]
    assert first["iec"] == "512345678"  # a number Excel stored loses no ".0"
    assert first["country"] == "India"
    assert parsed.rows[1][1]["industry"] == ""  # cells after the row's data are empty


def test_an_excel_file_is_recognised_by_its_content_whatever_its_name():
    content = _xlsx([list(TEMPLATE_COLUMNS), ["Acme", "IN"]])
    assert read_file("upload", content).rows[0][1]["name"] == "Acme"


def test_a_single_sheet_workbook_is_read_whatever_the_sheet_is_called():
    content = _xlsx([list(TEMPLATE_COLUMNS), ["Acme", "IN"]], sheet="Sheet1")
    assert len(read_file("companies.xlsx", content).rows) == 1


def test_a_workbook_with_several_sheets_needs_a_companies_sheet():
    content = _xlsx([list(TEMPLATE_COLUMNS)], sheet="Data", extra_sheet="Notes")
    with pytest.raises(ValidationError, match="no 'Companies' sheet"):
        read_file("companies.xlsx", content)


def test_cells_past_the_header_are_kept_apart():
    content = _xlsx([["name", "country"], ["Acme", "IN", "stray"]])
    [(_, raw)] = read_file("companies.xlsx", content).rows
    assert raw[None] == ["stray"]


def test_an_old_excel_file_is_refused_with_what_to_do():
    with pytest.raises(ValidationError, match=r"\.xls\)\. Save it as \.xlsx or \.csv"):
        read_file("companies.xls", b"\xd0\xcf\x11\xe0rest")


def test_a_file_named_xlsx_that_is_not_one_is_refused():
    with pytest.raises(ValidationError, match="not a readable Excel workbook"):
        read_file("companies.xlsx", b"name,country\nAcme,IN\n")


# ── Preview ───────────────────────────────────────────────────────────────────


def test_the_preview_shows_cells_in_their_columns_and_counts_every_row():
    rows = "\n".join(f"Co {i},IN,,,,,,," for i in range(15))
    preview = preview_file("companies.csv", (HEADER + "\n" + rows + "\n").encode())
    assert preview.columns == list(TEMPLATE_COLUMNS)
    assert preview.total_rows == 15
    assert len(preview.rows) == 10
    assert preview.rows[0].line == 2
    assert preview.rows[0].cells == ["Co 0", "IN", "", "", "", "", "", "", ""]
    assert preview.ready


def test_the_preview_reports_header_problems_instead_of_refusing():
    preview = preview_file("companies.csv", b"name,country,pan,actor_id\nAcme,IN,,x\n")
    assert preview.missing_columns == ["gstins", "iec", "cin", "source", "industry"]
    assert preview.unknown_columns == ["actor_id"]
    assert not preview.ready


def test_the_preview_reports_a_file_over_the_row_limit(monkeypatch):
    from app.modules.onboarding.application import company_import_service

    monkeypatch.setattr(company_import_service, "MAX_ROWS", 1)
    preview = preview_file("companies.csv", (HEADER + "\nA,IN\nB,IN\n").encode())
    assert (preview.total_rows, preview.max_rows, preview.ready) == (2, 1, False)


# ── Names for codes ───────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("value", "expected"),
    [("IN", "IN"), ("in", "IN"), (" India ", "IN"), ("united kingdom", "GB"),
     ("Atlantis", "Atlantis"), (None, None)],
)
def test_a_country_is_given_by_code_or_name(value, expected):
    assert country_code_for(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [(None, ExporterSource.MANUAL), ("REFERRAL", ExporterSource.REFERRAL),
     ("Existing customer", ExporterSource.EXISTING_CUSTOMER),
     ("existing_customer", ExporterSource.EXISTING_CUSTOMER),
     ("Manual entry", ExporterSource.MANUAL), ("RXIL", None), ("DEAL_BUYER", None),
     ("Cold call", None)],
)
def test_a_source_is_given_by_code_or_label_and_partner_sources_are_refused(value, expected):
    assert _source_for(value) is expected


# ── The Excel template ────────────────────────────────────────────────────────


def test_the_column_guide_describes_every_template_column_once():
    assert [guide.column for guide in COLUMN_GUIDE] == list(TEMPLATE_COLUMNS)


def test_the_excel_template_has_instructions_a_companies_sheet_and_dropdowns():
    workbook = load_workbook(io.BytesIO(build_template_xlsx()))
    assert workbook.sheetnames == ["Instructions", COMPANIES_SHEET, "Lists"]
    assert workbook["Lists"].sheet_state == "hidden"
    assert workbook.active.title == COMPANIES_SHEET

    companies = workbook[COMPANIES_SHEET]
    assert [c.value for c in companies[1]] == list(TEMPLATE_COLUMNS)
    validations = {str(v.sqref): v for v in companies.data_validations.dataValidation}
    assert set(validations) == {"B2:B1001", "G2:G1001"}  # country, source
    assert companies["C2"].number_format == "@"  # identifiers stay text

    lists = workbook["Lists"]
    assert len([c for c in lists["A"] if c.value]) == len(COUNTRY_NAMES)
    assert [c.value for c in lists["B"] if c.value] == list(IMPORT_SOURCE_LABELS.values())

    instructions = [c.value for c in workbook["Instructions"]["A"] if c.value]
    assert set(TEMPLATE_COLUMNS) <= set(instructions)


def test_the_excel_template_reads_back_as_an_empty_file_in_the_right_shape():
    preview = preview_file("company-import-template.xlsx", build_template_xlsx())
    assert preview.columns == list(TEMPLATE_COLUMNS)
    assert preview.total_rows == 0
    assert preview.ready
